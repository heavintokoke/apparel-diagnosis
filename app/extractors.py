from __future__ import annotations

import csv
import html
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

from PIL import Image
import openpyxl


WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


@dataclass
class ExtractedFile:
    filename: str
    file_type: str
    text: str
    tables: list[list[list[str]]]
    metadata: dict[str, Any]

    def to_dict(self) -> dict:
        return asdict(self)


def normalize_text(text: str) -> str:
    text = html.unescape(text or "")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _word_text(element: ET.Element) -> str:
    return "".join(node.text or "" for node in element.findall(".//w:t", WORD_NS)).strip()


def extract_docx(path: Path) -> ExtractedFile:
    paragraphs: list[str] = []
    tables: list[list[list[str]]] = []
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    for paragraph in root.findall(".//w:p", WORD_NS):
        text = _word_text(paragraph)
        if text:
            paragraphs.append(text)
    for table in root.findall(".//w:tbl", WORD_NS):
        rows: list[list[str]] = []
        for tr in table.findall("./w:tr", WORD_NS):
            row = [_word_text(tc) for tc in tr.findall("./w:tc", WORD_NS)]
            if any(row):
                rows.append(row)
        if rows:
            tables.append(rows)
    table_text = "\n".join(" | ".join(cell for cell in row) for table in tables for row in table)
    text = normalize_text("\n".join(paragraphs + [table_text]))
    return ExtractedFile(path.name, "Word", text, tables, {"paragraphs": len(paragraphs), "tables": len(tables)})


def extract_xlsx(path: Path) -> ExtractedFile:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    tables: list[list[list[str]]] = []
    text_lines: list[str] = []
    loaded_rows = 0
    numeric_cells = 0
    for sheet in workbook.worksheets[:20]:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(max_row=300, max_col=80, values_only=True):
            numeric_cells += sum(1 for value in row if isinstance(value, (int, float)) and not isinstance(value, bool))
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append(values)
        if rows:
            loaded_rows += len(rows)
            tables.append(rows)
            text_lines.append(f"工作表：{sheet.title}")
            text_lines.extend(" | ".join(row) for row in rows[:120])
    return ExtractedFile(
        path.name,
        "Excel",
        normalize_text("\n".join(text_lines)),
        tables,
        {"sheets": len(workbook.worksheets), "loaded_tables": len(tables), "loaded_rows": loaded_rows, "numeric_cells": numeric_cells},
    )


def _try_ocr_image(image: Image.Image) -> tuple[str, str]:
    tesseract_bin = shutil.which("tesseract")
    if not tesseract_bin:
        return "", "unavailable:tesseract_not_found"

    try:
        with tempfile.NamedTemporaryFile(suffix=".png") as temp_image:
            image.save(temp_image.name)
            completed = subprocess.run(
                [tesseract_bin, temp_image.name, "stdout", "-l", "chi_sim+eng"],
                check=False,
                capture_output=True,
                text=True,
                timeout=30,
            )
        if completed.returncode != 0:
            return "", f"failed:{completed.stderr.strip()[:160]}"
        text = completed.stdout
        text = normalize_text(text)
        return text, "applied" if text else "applied:no_text"
    except Exception as exc:
        return "", f"failed:{exc}"


def extract_pdf(path: Path) -> ExtractedFile:
    text_parts: list[str] = []
    tables: list[list[list[str]]] = []
    metadata: dict[str, Any] = {}
    low_text_pages = 0
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            metadata["pages"] = len(pdf.pages)
            for page in pdf.pages[:80]:
                page_text = page.extract_text() or ""
                if len(page_text.strip()) < 20:
                    low_text_pages += 1
                text_parts.append(page_text)
                for table in page.extract_tables() or []:
                    cleaned = [["" if cell is None else str(cell) for cell in row] for row in table]
                    if cleaned:
                        tables.append(cleaned)
    except Exception:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        metadata["pages"] = len(reader.pages)
        for page in reader.pages[:80]:
            page_text = page.extract_text() or ""
            if len(page_text.strip()) < 20:
                low_text_pages += 1
            text_parts.append(page_text)
    text = normalize_text("\n".join(text_parts))
    metadata["low_text_pages"] = low_text_pages
    if metadata.get("pages") and low_text_pages >= max(1, metadata["pages"] // 2) and len(text) < 200:
        metadata["scan_status"] = "likely_scanned_pdf_or_image_pdf"
        text = normalize_text(
            text
            + "\nPDF 可能是扫描件或图片型 PDF；当前环境未检测到可用 OCR 时，需要补充可复制文字版资料或安装 OCR 后重新分析。"
        )
    return ExtractedFile(path.name, "PDF", text, tables, metadata)


def extract_image(path: Path) -> ExtractedFile:
    with Image.open(path) as image:
        metadata = {
            "width": image.width,
            "height": image.height,
            "mode": image.mode,
            "format": image.format,
        }
        ocr_text, ocr_status = _try_ocr_image(image)
        metadata["ocr_status"] = ocr_status
    text = (
        f"图片文件：{path.name}\n"
        f"尺寸：{metadata['width']}x{metadata['height']}\n"
        + (f"OCR 识别文字：\n{ocr_text}" if ocr_text else "当前环境未完成 OCR 文字识别；如图片中包含表格或诊断资料，建议补充 Word/Excel/PDF 文字版。")
    )
    return ExtractedFile(path.name, "Image", text, [], metadata)


def extract_text_like(path: Path) -> ExtractedFile:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows: list[list[str]] = []
        with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as fh:
            for row in csv.reader(fh):
                if any(row):
                    rows.append(row)
        return ExtractedFile(path.name, "CSV", normalize_text("\n".join(" | ".join(row) for row in rows)), [rows], {})
    text = path.read_text(encoding="utf-8", errors="ignore")
    return ExtractedFile(path.name, "Text", normalize_text(text), [], {})


def extract_file(path: Path) -> ExtractedFile:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return extract_image(path)
    if suffix in {".txt", ".md", ".csv"}:
        return extract_text_like(path)
    raise ValueError(f"暂不支持的文件类型：{path.name}")


def build_information_package(paths: list[Path]) -> dict:
    extracted = [extract_file(path) for path in paths]
    combined = "\n\n".join(f"## {item.filename}\n{item.text}" for item in extracted if item.text)
    return {
        "files": [item.to_dict() for item in extracted],
        "combined_text": normalize_text(combined),
        "stats": {
            "file_count": len(extracted),
            "text_chars": len(combined),
            "tables": sum(len(item.tables) for item in extracted),
            "ocr_unavailable_files": sum(1 for item in extracted if str(item.metadata.get("ocr_status", "")).startswith("unavailable")),
            "scanned_pdf_candidates": sum(1 for item in extracted if item.metadata.get("scan_status") == "likely_scanned_pdf_or_image_pdf"),
        },
    }
